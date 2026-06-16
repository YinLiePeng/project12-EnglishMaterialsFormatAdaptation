#!/usr/bin/env python3
"""
PDF Parser Comparison Report Generator
Compares existing PyMuPDF parser vs opendataloader-pdf on 19 test samples.
"""

import json
import os
import sys
import time
from pathlib import Path
from typing import Dict, List, Any

# Add backend to path
sys.path.insert(0, str(Path(__file__).parent.parent))

import opendataloader_pdf
from docx import Document

# Paths
PROJECT_ROOT = Path("/mnt/e/Opencode/project12-EnglishMaterialsFormatAdaptation_2")
PDF_DIR = PROJECT_ROOT / "测试用例" / "原生PDF"
STANDARD_DIR = PROJECT_ROOT / "测试用例" / "网上下载的英语资料样例"
OUTPUT_DIR = PROJECT_ROOT / "backend" / "data" / "pdf_comparison"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


def get_pdf_files() -> List[Path]:
    """Get all PDF files from test directory."""
    files = sorted(PDF_DIR.glob("*.pdf"))
    return files


def analyze_standard_docx(docx_path: Path) -> Dict[str, Any]:
    """Analyze standard DOCX to get ground truth metrics."""
    try:
        doc = Document(str(docx_path))

        paragraphs = 0
        text_paragraphs = 0
        blank_paragraphs = 0
        tables = len(doc.tables)
        total_chars = 0

        for para in doc.paragraphs:
            paragraphs += 1
            text = para.text.strip()
            if text:
                text_paragraphs += 1
                total_chars += len(text)
            else:
                blank_paragraphs += 1

        # Count images (inline shapes in document)
        images = 0
        try:
            # Count relationships that are images
            rels = doc.part.rels
            for rel in rels.values():
                if "image" in rel.reltype:
                    images += 1
        except:
            pass

        return {
            "paragraphs": paragraphs,
            "text_paragraphs": text_paragraphs,
            "blank_paragraphs": blank_paragraphs,
            "tables": tables,
            "images": images,
            "total_chars": total_chars,
            "avg_para_length": round(total_chars / text_paragraphs, 1)
            if text_paragraphs > 0
            else 0,
        }
    except Exception as e:
        return {"error": str(e)}


def analyze_with_existing_parser(pdf_path: Path) -> Dict[str, Any]:
    """Analyze PDF using existing PyMuPDF parser."""
    try:
        from app.services.pdf.parser import PDFParser

        start_time = time.time()
        with PDFParser(str(pdf_path)) as parser:
            content_list = parser.extract_structured_content()
        elapsed = time.time() - start_time

        paragraphs = len(content_list)
        text_paragraphs = 0
        blank_paragraphs = 0
        tables = 0
        images = 0
        total_chars = 0

        for item in content_list:
            if item.get("type") == "TABLE":
                tables += 1
            elif item.get("type") == "IMAGE":
                images += 1
            elif item.get("type") == "BLANK_LINE":
                blank_paragraphs += 1
            elif item.get("text"):
                text = item.get("text", "").strip()
                if text:
                    text_paragraphs += 1
                    total_chars += len(text)
                else:
                    blank_paragraphs += 1
            else:
                blank_paragraphs += 1

        return {
            "paragraphs": paragraphs,
            "text_paragraphs": text_paragraphs,
            "blank_paragraphs": blank_paragraphs,
            "tables": tables,
            "images": images,
            "total_chars": total_chars,
            "avg_para_length": round(total_chars / text_paragraphs, 1)
            if text_paragraphs > 0
            else 0,
            "elapsed_time": round(elapsed, 2),
        }
    except Exception as e:
        return {"error": str(e)}


def analyze_with_opendataloader(pdf_path: Path) -> Dict[str, Any]:
    """Analyze PDF using opendataloader-pdf."""
    try:
        start_time = time.time()
        opendataloader_pdf.convert(
            input_path=str(pdf_path),
            output_dir=str(OUTPUT_DIR),
            format="json",
            include_header_footer=False,
        )
        elapsed = time.time() - start_time

        # Read the generated JSON
        json_file = OUTPUT_DIR / f"{pdf_path.stem}.json"
        if not json_file.exists():
            json_file = OUTPUT_DIR / f"{pdf_path.stem}_opendataloader.json"

        if not json_file.exists():
            return {
                "error": f"JSON output not found: {list(OUTPUT_DIR.glob('*.json'))}"
            }

        with open(json_file, "r", encoding="utf-8") as f:
            data = json.load(f)

        # Recursive function to traverse nested kids structure
        def traverse(node):
            counts = {
                "paragraphs": 0,
                "text_paragraphs": 0,
                "blank_paragraphs": 0,
                "tables": 0,
                "images": 0,
                "headings": 0,
                "lists": 0,
                "total_chars": 0,
            }

            if isinstance(node, dict):
                elem_type = node.get("type", "").lower()
                content = node.get("content", "") or ""

                if elem_type == "table":
                    counts["tables"] += 1
                elif elem_type == "image":
                    counts["images"] += 1
                elif elem_type == "heading":
                    counts["headings"] += 1
                    counts["paragraphs"] += 1
                    text = str(content).strip()
                    if text:
                        counts["text_paragraphs"] += 1
                        counts["total_chars"] += len(text)
                elif elem_type == "list":
                    counts["lists"] += 1
                elif elem_type == "paragraph":
                    counts["paragraphs"] += 1
                    text = str(content).strip()
                    if text:
                        counts["text_paragraphs"] += 1
                        counts["total_chars"] += len(text)
                    else:
                        counts["blank_paragraphs"] += 1

                # Recurse into kids
                for child in node.get("kids", []):
                    child_counts = traverse(child)
                    for k in counts:
                        counts[k] += child_counts[k]

            elif isinstance(node, list):
                for item in node:
                    child_counts = traverse(item)
                    for k in counts:
                        counts[k] += child_counts[k]

            return counts

        # Start traversal from top-level kids
        elements = (
            data.get("kids", [])
            if isinstance(data, dict)
            else (data if isinstance(data, list) else [])
        )
        result = traverse(elements)

        return {
            "paragraphs": result["paragraphs"],
            "text_paragraphs": result["text_paragraphs"],
            "blank_paragraphs": result["blank_paragraphs"],
            "tables": result["tables"],
            "images": result["images"],
            "headings": result["headings"],
            "lists": result["lists"],
            "total_chars": result["total_chars"],
            "avg_para_length": round(
                result["total_chars"] / result["text_paragraphs"], 1
            )
            if result["text_paragraphs"] > 0
            else 0,
            "elapsed_time": round(elapsed, 2),
        }
    except Exception as e:
        import traceback

        return {"error": f"{str(e)}\n{traceback.format_exc()}"}


def generate_html_report(results: List[Dict[str, Any]]) -> str:
    """Generate HTML comparison report."""

    html = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>PDF Parser 对比报告</title>
    <style>
        * { margin: 0; padding: 0; box-sizing: border-box; }
        body { font-family: -apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, sans-serif; background: #f5f5f5; padding: 20px; color: #333; }
        .container { max-width: 1400px; margin: 0 auto; }
        h1 { text-align: center; color: #1a1a1a; margin-bottom: 10px; font-size: 28px; }
        .subtitle { text-align: center; color: #666; margin-bottom: 30px; font-size: 14px; }
        .summary { background: white; border-radius: 12px; padding: 24px; margin-bottom: 24px; box-shadow: 0 2px 8px rgba(0,0,0,0.06); }
        .summary h2 { font-size: 18px; margin-bottom: 16px; color: #1a1a1a; }
        .summary-grid { display: grid; grid-template-columns: repeat(auto-fit, minmax(200px, 1fr)); gap: 16px; }
        .summary-card { background: #f8f9fa; border-radius: 8px; padding: 16px; text-align: center; }
        .summary-card .label { font-size: 12px; color: #666; margin-bottom: 4px; }
        .summary-card .value { font-size: 24px; font-weight: 700; color: #1a1a1a; }
        .summary-card .sub { font-size: 11px; color: #999; margin-top: 4px; }
        .file-card { background: white; border-radius: 12px; padding: 20px; margin-bottom: 16px; box-shadow: 0 2px 8px rgba(0,0,0,0.06); }
        .file-header { display: flex; justify-content: space-between; align-items: center; margin-bottom: 16px; padding-bottom: 12px; border-bottom: 1px solid #eee; }
        .file-name { font-size: 15px; font-weight: 600; color: #1a1a1a; max-width: 70%; overflow: hidden; text-overflow: ellipsis; white-space: nowrap; }
        .file-status { font-size: 12px; padding: 4px 10px; border-radius: 20px; font-weight: 500; }
        .status-ok { background: #e8f5e9; color: #2e7d32; }
        .status-warn { background: #fff3e0; color: #ef6c00; }
        .status-error { background: #ffebee; color: #c62828; }
        .metrics-grid { display: grid; grid-template-columns: repeat(4, 1fr); gap: 12px; }
        .metric-group { background: #f8f9fa; border-radius: 8px; padding: 12px; }
        .metric-group h4 { font-size: 12px; color: #666; margin-bottom: 8px; font-weight: 500; text-transform: uppercase; letter-spacing: 0.5px; }
        .metric-row { display: flex; justify-content: space-between; align-items: center; padding: 4px 0; font-size: 13px; border-bottom: 1px solid #eee; }
        .metric-row:last-child { border-bottom: none; }
        .metric-label { color: #666; }
        .metric-value { font-weight: 600; font-family: 'SF Mono', Monaco, monospace; }
        .diff-positive { color: #2e7d32; }
        .diff-negative { color: #c62828; }
        .diff-neutral { color: #666; }
        .legend { display: flex; gap: 16px; justify-content: center; margin-bottom: 20px; font-size: 12px; }
        .legend-item { display: flex; align-items: center; gap: 4px; }
        .legend-dot { width: 10px; height: 10px; border-radius: 50%; }
        .dot-standard { background: #2196f3; }
        .dot-existing { background: #ff9800; }
        .dot-opendataloader { background: #4caf50; }
        .progress-bar { width: 100%; height: 4px; background: #e0e0e0; border-radius: 2px; margin-top: 8px; overflow: hidden; }
        .progress-fill { height: 100%; border-radius: 2px; transition: width 0.3s; }
        .progress-standard { background: #2196f3; }
        .progress-existing { background: #ff9800; }
        .progress-opendataloader { background: #4caf50; }
        table { width: 100%; border-collapse: collapse; font-size: 13px; }
        th { background: #f5f5f5; padding: 10px; text-align: left; font-weight: 600; color: #555; border-bottom: 2px solid #ddd; }
        td { padding: 10px; border-bottom: 1px solid #eee; }
        tr:hover { background: #fafafa; }
        .num { font-family: 'SF Mono', Monaco, monospace; text-align: right; }
        .best { background: #e8f5e9 !important; font-weight: 600; }
        .worst { background: #ffebee !important; }
        .section { margin-bottom: 32px; }
        .section h2 { font-size: 18px; margin-bottom: 16px; color: #1a1a1a; display: flex; align-items: center; gap: 8px; }
        .badge { display: inline-block; padding: 2px 8px; border-radius: 4px; font-size: 11px; font-weight: 600; }
        .badge-green { background: #e8f5e9; color: #2e7d32; }
        .badge-orange { background: #fff3e0; color: #ef6c00; }
        .badge-red { background: #ffebee; color: #c62828; }
    </style>
</head>
<body>
    <div class="container">
        <h1>PDF Parser 对比报告</h1>
        <p class="subtitle">现有 PyMuPDF Parser vs OpenDataLoader PDF · 19 个原生 PDF 样本</p>
        
        <div class="legend">
            <div class="legend-item"><div class="legend-dot dot-standard"></div>标准 DOCX (Ground Truth)</div>
            <div class="legend-item"><div class="legend-dot dot-existing"></div>现有 Parser</div>
            <div class="legend-item"><div class="legend-dot dot-opendataloader"></div>OpenDataLoader</div>
        </div>
"""

    # Calculate summary statistics
    total_files = len(results)
    success_existing = sum(1 for r in results if "error" not in r.get("existing", {}))
    success_opendataloader = sum(
        1 for r in results if "error" not in r.get("opendataloader", {})
    )

    # Calculate average deviations
    para_dev_existing = []
    para_dev_opendataloader = []
    table_dev_existing = []
    table_dev_opendataloader = []
    char_dev_existing = []
    char_dev_opendataloader = []

    for r in results:
        std = r.get("standard", {})
        ex = r.get("existing", {})
        od = r.get("opendataloader", {})

        if "error" not in std and "error" not in ex:
            std_para = std.get("paragraphs", 0)
            if std_para > 0:
                para_dev_existing.append(
                    abs(ex.get("paragraphs", 0) - std_para) / std_para
                )
                char_dev_existing.append(
                    abs(ex.get("total_chars", 0) - std.get("total_chars", 0))
                    / max(std.get("total_chars", 1), 1)
                )

        if "error" not in std and "error" not in od:
            std_para = std.get("paragraphs", 0)
            if std_para > 0:
                para_dev_opendataloader.append(
                    abs(od.get("paragraphs", 0) - std_para) / std_para
                )
                char_dev_opendataloader.append(
                    abs(od.get("total_chars", 0) - std.get("total_chars", 0))
                    / max(std.get("total_chars", 1), 1)
                )

    avg_para_dev_ex = (
        round(sum(para_dev_existing) / len(para_dev_existing) * 100, 1)
        if para_dev_existing
        else 0
    )
    avg_para_dev_od = (
        round(sum(para_dev_opendataloader) / len(para_dev_opendataloader) * 100, 1)
        if para_dev_opendataloader
        else 0
    )
    avg_char_dev_ex = (
        round(sum(char_dev_existing) / len(char_dev_existing) * 100, 1)
        if char_dev_existing
        else 0
    )
    avg_char_dev_od = (
        round(sum(char_dev_opendataloader) / len(char_dev_opendataloader) * 100, 1)
        if char_dev_opendataloader
        else 0
    )

    # Determine winner
    winner = "opendataloader" if avg_para_dev_od < avg_para_dev_ex else "existing"
    winner_para = min(avg_para_dev_ex, avg_para_dev_od)

    html += f"""
        <div class="summary">
            <h2>总体概览</h2>
            <div class="summary-grid">
                <div class="summary-card">
                    <div class="label">测试文件数</div>
                    <div class="value">{total_files}</div>
                </div>
                <div class="summary-card">
                    <div class="label">现有 Parser 成功率</div>
                    <div class="value" style="color: {"#2e7d32" if success_existing == total_files else "#ef6c00"}">{success_existing}/{total_files}</div>
                </div>
                <div class="summary-card">
                    <div class="label">OpenDataLoader 成功率</div>
                    <div class="value" style="color: {"#2e7d32" if success_opendataloader == total_files else "#ef6c00"}">{success_opendataloader}/{total_files}</div>
                </div>
                <div class="summary-card">
                    <div class="label">段落偏差 (现有)</div>
                    <div class="value" style="color: {"#2e7d32" if avg_para_dev_ex < 50 else "#ef6c00" if avg_para_dev_ex < 100 else "#c62828"}">{avg_para_dev_ex}%</div>
                    <div class="sub">平均偏离标准 DOCX</div>
                </div>
                <div class="summary-card">
                    <div class="label">段落偏差 (OpenDataLoader)</div>
                    <div class="value" style="color: {"#2e7d32" if avg_para_dev_od < 50 else "#ef6c00" if avg_para_dev_od < 100 else "#c62828"}">{avg_para_dev_od}%</div>
                    <div class="sub">平均偏离标准 DOCX</div>
                </div>
                <div class="summary-card">
                    <div class="label">字符偏差 (现有)</div>
                    <div class="value">{avg_char_dev_ex}%</div>
                </div>
                <div class="summary-card">
                    <div class="label">字符偏差 (OpenDataLoader)</div>
                    <div class="value">{avg_char_dev_od}%</div>
                </div>
                <div class="summary-card">
                    <div class="label">段落还原 winner</div>
                    <div class="value" style="color: #2e7d32">{winner.upper()}</div>
                    <div class="sub">偏差 {winner_para}%</div>
                </div>
            </div>
        </div>
"""

    # Detailed per-file cards
    html += '<div class="section"><h2>逐个文件详细对比</h2></div>'

    for i, r in enumerate(results, 1):
        filename = r["filename"]
        std = r.get("standard", {})
        ex = r.get("existing", {})
        od = r.get("opendataloader", {})

        ex_error = "error" in ex
        od_error = "error" in od

        # Determine status
        if ex_error and od_error:
            status_class = "status-error"
            status_text = "两者都失败"
        elif ex_error:
            status_class = "status-warn"
            status_text = "现有 Parser 失败"
        elif od_error:
            status_class = "status-warn"
            status_text = "OpenDataLoader 失败"
        else:
            status_class = "status-ok"
            status_text = "正常"

        html += f"""
        <div class="file-card">
            <div class="file-header">
                <div class="file-name">{i}. {filename}</div>
                <div class="file-status {status_class}">{status_text}</div>
            </div>
            <div class="metrics-grid">
                <div class="metric-group">
                    <h4>标准 DOCX</h4>
                    <div class="metric-row"><span class="metric-label">段落</span><span class="metric-value num">{
            std.get("paragraphs", "N/A")
        }</span></div>
                    <div class="metric-row"><span class="metric-label">文本段落</span><span class="metric-value num">{
            std.get("text_paragraphs", "N/A")
        }</span></div>
                    <div class="metric-row"><span class="metric-label">空白段落</span><span class="metric-value num">{
            std.get("blank_paragraphs", "N/A")
        }</span></div>
                    <div class="metric-row"><span class="metric-label">表格</span><span class="metric-value num">{
            std.get("tables", "N/A")
        }</span></div>
                    <div class="metric-row"><span class="metric-label">图片</span><span class="metric-value num">{
            std.get("images", "N/A")
        }</span></div>
                    <div class="metric-row"><span class="metric-label">总字符</span><span class="metric-value num">{
            std.get("total_chars", "N/A")
        }</span></div>
                </div>
                <div class="metric-group">
                    <h4>现有 Parser</h4>
                    {
            "<div class='metric-row'><span class='metric-label'>错误</span><span class='metric-value' style='color:#c62828'>"
            + ex.get("error", "")[:50]
            + "</span></div>"
            if ex_error
            else f'''
                    <div class="metric-row"><span class="metric-label">段落</span><span class="metric-value num">{ex.get("paragraphs", 0)}</span></div>
                    <div class="metric-row"><span class="metric-label">文本段落</span><span class="metric-value num">{ex.get("text_paragraphs", 0)}</span></div>
                    <div class="metric-row"><span class="metric-label">空白段落</span><span class="metric-value num">{ex.get("blank_paragraphs", 0)}</span></div>
                    <div class="metric-row"><span class="metric-label">表格</span><span class="metric-value num">{ex.get("tables", 0)}</span></div>
                    <div class="metric-row"><span class="metric-label">图片</span><span class="metric-value num">{ex.get("images", 0)}</span></div>
                    <div class="metric-row"><span class="metric-label">总字符</span><span class="metric-value num">{ex.get("total_chars", 0)}</span></div>
                    <div class="metric-row"><span class="metric-label">耗时</span><span class="metric-value num">{ex.get("elapsed_time", 0)}s</span></div>
                    '''
        }
                </div>
                <div class="metric-group">
                    <h4>OpenDataLoader</h4>
                    {
            "<div class='metric-row'><span class='metric-label'>错误</span><span class='metric-value' style='color:#c62828'>"
            + od.get("error", "")[:50]
            + "</span></div>"
            if od_error
            else f'''
                    <div class="metric-row"><span class="metric-label">段落</span><span class="metric-value num">{od.get("paragraphs", 0)}</span></div>
                    <div class="metric-row"><span class="metric-label">文本段落</span><span class="metric-value num">{od.get("text_paragraphs", 0)}</span></div>
                    <div class="metric-row"><span class="metric-label">空白段落</span><span class="metric-value num">{od.get("blank_paragraphs", 0)}</span></div>
                    <div class="metric-row"><span class="metric-label">表格</span><span class="metric-value num">{od.get("tables", 0)}</span></div>
                    <div class="metric-row"><span class="metric-label">图片</span><span class="metric-value num">{od.get("images", 0)}</span></div>
                    <div class="metric-row"><span class="metric-label">总字符</span><span class="metric-value num">{od.get("total_chars", 0)}</span></div>
                    <div class="metric-row"><span class="metric-label">耗时</span><span class="metric-value num">{od.get("elapsed_time", 0)}s</span></div>
                    '''
        }
                </div>
                <div class="metric-group">
                    <h4>偏离度对比</h4>
                    {
            "<div class='metric-row'><span class='metric-label'>无法计算</span><span class='metric-value'>数据缺失</span></div>"
            if (ex_error or od_error or "error" in std)
            else f'''
                    <div class="metric-row">
                        <span class="metric-label">段落偏离</span>
                        <span class="metric-value">
                            <span style="color: #ff9800">现有: {round(abs(ex.get("paragraphs", 0) - std.get("paragraphs", 0)) / max(std.get("paragraphs", 1), 1) * 100, 1)}%</span><br>
                            <span style="color: #4caf50">OD: {round(abs(od.get("paragraphs", 0) - std.get("paragraphs", 0)) / max(std.get("paragraphs", 1), 1) * 100, 1)}%</span>
                        </span>
                    </div>
                    <div class="metric-row">
                        <span class="metric-label">字符偏离</span>
                        <span class="metric-value">
                            <span style="color: #ff9800">现有: {round(abs(ex.get("total_chars", 0) - std.get("total_chars", 0)) / max(std.get("total_chars", 1), 1) * 100, 1)}%</span><br>
                            <span style="color: #4caf50">OD: {round(abs(od.get("total_chars", 0) - std.get("total_chars", 0)) / max(std.get("total_chars", 1), 1) * 100, 1)}%</span>
                        </span>
                    </div>
                    <div class="metric-row">
                        <span class="metric-label">表格偏离</span>
                        <span class="metric-value">
                            <span style="color: #ff9800">现有: {abs(ex.get("tables", 0) - std.get("tables", 0))}</span><br>
                            <span style="color: #4caf50">OD: {abs(od.get("tables", 0) - std.get("tables", 0))}</span>
                        </span>
                    </div>
                    <div class="metric-row">
                        <span class="metric-label">图片偏离</span>
                        <span class="metric-value">
                            <span style="color: #ff9800">现有: {abs(ex.get("images", 0) - std.get("images", 0))}</span><br>
                            <span style="color: #4caf50">OD: {abs(od.get("images", 0) - std.get("images", 0))}</span>
                        </span>
                    </div>
                    '''
        }
                </div>
            </div>
        </div>
"""

    # Summary table
    html += """
        <div class="section">
            <h2>汇总数据表</h2>
            <table>
                <thead>
                    <tr>
                        <th>#</th>
                        <th>文件名</th>
                        <th class="num">标准段落</th>
                        <th class="num">现有段落</th>
                        <th class="num">OD 段落</th>
                        <th class="num">标准表格</th>
                        <th class="num">现有表格</th>
                        <th class="num">OD 表格</th>
                        <th class="num">标准字符</th>
                        <th class="num">现有字符</th>
                        <th class="num">OD 字符</th>
                        <th>段落偏差 Winner</th>
                    </tr>
                </thead>
                <tbody>
"""

    for i, r in enumerate(results, 1):
        filename = (
            r["filename"][:40] + "..." if len(r["filename"]) > 40 else r["filename"]
        )
        std = r.get("standard", {})
        ex = r.get("existing", {})
        od = r.get("opendataloader", {})

        std_para = std.get("paragraphs", 0)
        ex_para = ex.get("paragraphs", 0) if "error" not in ex else "ERR"
        od_para = od.get("paragraphs", 0) if "error" not in od else "ERR"

        std_tbl = std.get("tables", 0)
        ex_tbl = ex.get("tables", 0) if "error" not in ex else "ERR"
        od_tbl = od.get("tables", 0) if "error" not in od else "ERR"

        std_char = std.get("total_chars", 0)
        ex_char = ex.get("total_chars", 0) if "error" not in ex else "ERR"
        od_char = od.get("total_chars", 0) if "error" not in od else "ERR"

        # Determine winner
        if "error" not in ex and "error" not in od and std_para > 0:
            ex_dev = abs(ex_para - std_para) / std_para
            od_dev = abs(od_para - std_para) / std_para
            winner_text = "OpenDataLoader" if od_dev < ex_dev else "现有 Parser"
            winner_class = "best" if od_dev < ex_dev else "worst"
        else:
            winner_text = "N/A"
            winner_class = ""

        html += f"""
                    <tr>
                        <td>{i}</td>
                        <td>{filename}</td>
                        <td class="num">{std_para}</td>
                        <td class="num">{ex_para}</td>
                        <td class="num">{od_para}</td>
                        <td class="num">{std_tbl}</td>
                        <td class="num">{ex_tbl}</td>
                        <td class="num">{od_tbl}</td>
                        <td class="num">{std_char}</td>
                        <td class="num">{ex_char}</td>
                        <td class="num">{od_char}</td>
                        <td class="{winner_class}">{winner_text}</td>
                    </tr>
"""

    html += (
        """
                </tbody>
            </table>
        </div>
        
        <div style="text-align: center; color: #999; font-size: 12px; margin-top: 40px; padding-bottom: 20px;">
            生成时间: """
        + time.strftime("%Y-%m-%d %H:%M:%S")
        + """ | 分支: pdf-comparison
        </div>
    </div>
</body>
</html>
"""
    )

    return html


def main():
    pdf_files = get_pdf_files()
    print(f"Found {len(pdf_files)} PDF files to compare")

    results = []

    for i, pdf_file in enumerate(pdf_files, 1):
        print(f"\n[{i}/{len(pdf_files)}] Processing: {pdf_file.name}")

        # Find corresponding standard DOCX
        docx_name = pdf_file.stem + ".docx"
        docx_path = STANDARD_DIR / docx_name

        result = {"filename": pdf_file.name}

        # Analyze standard DOCX
        if docx_path.exists():
            print(f"  Standard DOCX found: {docx_path.name}")
            result["standard"] = analyze_standard_docx(docx_path)
        else:
            print(f"  WARNING: Standard DOCX not found: {docx_name}")
            result["standard"] = {"error": "Not found"}

        # Analyze with existing parser
        print(f"  Running existing parser...")
        result["existing"] = analyze_with_existing_parser(pdf_file)
        if "error" in result["existing"]:
            print(f"  ERROR (existing): {result['existing']['error'][:100]}")
        else:
            print(
                f"  Existing: {result['existing'].get('paragraphs', 0)} paragraphs, {result['existing'].get('elapsed_time', 0)}s"
            )

        # Analyze with opendataloader
        print(f"  Running opendataloader...")
        result["opendataloader"] = analyze_with_opendataloader(pdf_file)
        if "error" in result["opendataloader"]:
            print(
                f"  ERROR (opendataloader): {result['opendataloader']['error'][:100]}"
            )
        else:
            print(
                f"  OpenDataLoader: {result['opendataloader'].get('paragraphs', 0)} paragraphs, {result['opendataloader'].get('elapsed_time', 0)}s"
            )

        results.append(result)

    # Save raw results
    results_file = OUTPUT_DIR / "comparison_results.json"
    with open(results_file, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(f"\nRaw results saved to: {results_file}")

    # Generate HTML report
    html = generate_html_report(results)
    report_file = OUTPUT_DIR / "comparison_report.html"
    with open(report_file, "w", encoding="utf-8") as f:
        f.write(html)
    print(f"HTML report saved to: {report_file}")

    print("\nDone!")


if __name__ == "__main__":
    main()
