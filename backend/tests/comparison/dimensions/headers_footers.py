"""页眉页脚对比维度"""

from difflib import SequenceMatcher
from typing import List, Tuple, Dict, Any
from pathlib import Path

from .base import ComparisonDimension


class HeadersFootersDimension(ComparisonDimension):
    """对比页眉页脚：文本内容、链接状态"""

    name = "headers_footers"
    weight = 0.05
    description = "对比页眉页脚内容和格式一致性"

    def compare(
        self,
        gen_elements: List,
        ref_elements: List,
        gen_path: Path,
        ref_path: Path,
        match_result: Any = None,
    ) -> Tuple[float, Dict[str, Any]]:
        from docx import Document as DocxDocument

        gen_sections = self._get_sections(gen_path, DocxDocument)
        ref_sections = self._get_sections(ref_path, DocxDocument)

        if gen_sections is None and ref_sections is None:
            return 1.0, {
                "header_text_similarity": 1.0,
                "footer_text_similarity": 1.0,
                "presence_match": 1.0,
                "gen_section_count": 0,
                "ref_section_count": 0,
            }

        if gen_sections is None or ref_sections is None:
            return 0.0, {
                "header_text_similarity": 0.0,
                "footer_text_similarity": 0.0,
                "presence_match": 0.0,
                "gen_section_count": 0 if gen_sections is None else len(gen_sections),
                "ref_section_count": 0 if ref_sections is None else len(ref_sections),
            }

        header_scores = []
        footer_scores = []
        presence_scores = []
        details_per_section = []

        max_sections = max(len(gen_sections), len(ref_sections))

        for i in range(max_sections):
            gen_sec = gen_sections[i] if i < len(gen_sections) else None
            ref_sec = ref_sections[i] if i < len(ref_sections) else None

            # Header comparison
            h_score, h_detail = self._compare_header_footer(
                gen_sec, ref_sec, "header"
            )
            header_scores.append(h_score)

            # Footer comparison
            f_score, f_detail = self._compare_header_footer(
                gen_sec, ref_sec, "footer"
            )
            footer_scores.append(f_score)

            # Presence match
            gen_has_h = gen_sec is not None and self._has_content(gen_sec, "header")
            ref_has_h = ref_sec is not None and self._has_content(ref_sec, "header")
            gen_has_f = gen_sec is not None and self._has_content(gen_sec, "footer")
            ref_has_f = ref_sec is not None and self._has_content(ref_sec, "footer")

            presence = 0.0
            h_presence = 1.0 if gen_has_h == ref_has_h else 0.0
            f_presence = 1.0 if gen_has_f == ref_has_f else 0.0
            presence = (h_presence + f_presence) / 2.0
            presence_scores.append(presence)

            details_per_section.append({
                "section_index": i,
                "header_similarity": round(h_score, 4),
                "footer_similarity": round(f_score, 4),
                "presence_match": round(presence, 4),
            })

        header_text_sim = self._mean(header_scores)
        footer_text_sim = self._mean(footer_scores)
        presence_match = self._mean(presence_scores)

        score = header_text_sim * 0.35 + footer_text_sim * 0.35 + presence_match * 0.30
        score = max(0.0, min(1.0, score))

        return score, {
            "header_text_similarity": round(header_text_sim, 4),
            "footer_text_similarity": round(footer_text_sim, 4),
            "presence_match": round(presence_match, 4),
            "gen_section_count": len(gen_sections),
            "ref_section_count": len(ref_sections),
            "sections": details_per_section,
        }

    @staticmethod
    def _get_sections(path: Path, DocxDocument) -> Any:
        """读取文档的 sections"""
        try:
            doc = DocxDocument(str(path))
            return list(doc.sections)
        except Exception:
            return None

    @staticmethod
    def _has_content(section, part: str) -> bool:
        """检查 section 的 header/footer 是否有内容"""
        try:
            if part == "header":
                hf = section.header
            else:
                hf = section.footer

            if hf is None:
                return False
            if hf.is_linked_to_previous:
                return False

            text = hf.paragraphs[0].text.strip() if hf.paragraphs else ""
            return len(text) > 0
        except Exception:
            return False

    @staticmethod
    def _get_text(section, part: str) -> str:
        """获取 header/footer 的文本内容"""
        try:
            if part == "header":
                hf = section.header
            else:
                hf = section.footer

            if hf is None or hf.is_linked_to_previous:
                return ""

            texts = []
            for para in hf.paragraphs:
                t = para.text.strip()
                if t:
                    texts.append(t)
            return "\n".join(texts)
        except Exception:
            return ""

    @staticmethod
    def _is_linked_to_previous(section, part: str) -> bool:
        """检查 header/footer 是否链接到上一节"""
        try:
            if part == "header":
                hf = section.header
            else:
                hf = section.footer
            return hf.is_linked_to_previous
        except Exception:
            return False

    def _compare_header_footer(
        self, gen_sec, ref_sec, part: str
    ) -> Tuple[float, Dict[str, Any]]:
        """比较一个 section 的 header 或 footer"""
        if gen_sec is None and ref_sec is None:
            return 1.0, {}

        if gen_sec is None or ref_sec is None:
            return 0.0, {}

        gen_text = self._get_text(gen_sec, part)
        ref_text = self._get_text(ref_sec, part)

        gen_linked = self._is_linked_to_previous(gen_sec, part)
        ref_linked = self._is_linked_to_previous(ref_sec, part)

        # 文本相似度
        if not gen_text and not ref_text:
            text_sim = 1.0
        elif not gen_text or not ref_text:
            text_sim = 0.0
        else:
            text_sim = SequenceMatcher(None, gen_text, ref_text).ratio()

        # 链接状态匹配
        link_match = 1.0 if gen_linked == ref_linked else 0.0

        score = text_sim * 0.7 + link_match * 0.3

        detail = {
            f"{part}_gen_text_len": len(gen_text),
            f"{part}_ref_text_len": len(ref_text),
            f"{part}_text_similarity": round(text_sim, 4),
            f"{part}_gen_linked": gen_linked,
            f"{part}_ref_linked": ref_linked,
            f"{part}_link_match": link_match == 1.0,
        }

        return score, detail
